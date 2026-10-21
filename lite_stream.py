# lite_stream.py
# ─────────────────────────────────────────────────────────────────────────────
# Sends plain chat turns to the lightweight Gemini streaming service
# (zindryx-ai-backend /stream) instead of running the full agent for them.
#
# Env vars (set on THIS backend in Render):
#   LITE_STREAM_URL   e.g. https://zindryx-ai-backend.onrender.com/stream
#                     (leave unset to turn the whole feature OFF)
#   SERVICE_API_KEY   shared secret; the SAME value is set on the Gemini service
#
# Events yielded use the SAME shapes main.py already sends to the app:
#   {"type": "text",  "delta": "..."}
#   {"type": "block", "kind": "images", "images": [...]}
#   {"type": "final", "answer": ..., "segments": [...], "blocks": [...], "images": [...], ...}
#
# If the service can't be reached BEFORE it sends anything, LiteUnavailable is
# raised so main.py can fall back to the normal full agent. Nothing is lost.
# ─────────────────────────────────────────────────────────────────────────────
import io
import json
import os
import re
from typing import Iterator, List, Optional

import requests

from gpt2_tools import build_answer_segments

LITE_STREAM_URL = (os.getenv("LITE_STREAM_URL") or "").strip()
SERVICE_API_KEY = os.getenv("SERVICE_API_KEY") or ""

# (connect timeout, max silence between chunks). Free-tier Render can be asleep;
# if it isn't answering in time we fall back to the normal agent instead of making the user wait.
TIMEOUT = (10, 30)

HISTORY_MESSAGES = 6          # only the last few turns are sent (the service trims again)
HISTORY_CHARS = 700

# Things the lightweight service can't do -> use the full agent (files, code, live web, image GENERATION...).
# Edit freely. Image *search* ("show me pictures of ...") is NOT in this list: the lite service handles it.
_HEAVY = re.compile(
    # CREATING / exporting files (reading an attached pdf/docx is fine on the lite path)
    r"\b(make|create|generate|build|export|convert|save|write|turn|give me|send me)\b[^.?!\n]{0,40}\b(pdf|docx?|word file|xlsx?|excel|powerpoint|pptx|slides?|spreadsheet|csv|zip)\b"
    r"|\b(download|run (the |this )?code|python script|write (me )?(a )?(code|program|script)|"
    r"build (me )?(a |an )?(app|website|web ?page|game)|generate (an? )?image|create (an? )?image|"
    r"draw (me )?(a|an)\b|make (an? )?(logo|poster)|"
    r"latest|news|today|tonight|right now|current (news|affairs|events|price)|trending|price of|weather|"
    r"search (the )?(web|internet|online)|google|browse|workspace|github|video|audio|voice)\b",
    re.I,
)


class LiteUnavailable(Exception):
    pass


def enabled() -> bool:
    return bool(LITE_STREAM_URL)


_IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".heic"}
_DOC_EXT = {".pdf", ".docx", ".xlsx", ".xlsm"}
_TEXT_EXT = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".css", ".js", ".ts", ".py", ".dart", ".java",
             ".c", ".cpp", ".xml", ".yaml", ".yml", ".log", ".sql", ".kt", ".php", ".tsx", ".jsx"}
MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_FILE_TEXT = 15000


def _ext(url: str, name: str = "") -> str:
    return os.path.splitext((name or url.split("?")[0]).lower())[1]


def should_use_lite(question: str, image_urls, file_urls, file_names=None) -> bool:
    """Chat turns the lite service can handle: text, images, and common documents.
    Zips, slides, audio/video etc. -> full agent (it has the workspace + tools for those)."""
    if not enabled() or _HEAVY.search(question or ""):
        return False
    names = list(file_names or [])
    for i, u in enumerate(file_urls or []):
        e = _ext(u, names[i] if i < len(names) else "")
        if e not in _DOC_EXT | _TEXT_EXT | _IMG_EXT:
            return False
    return True


# ── attachment reading ──────────────────────────────────────────────────────
def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(f"[page {i}]\n{(pg.extract_text() or '').strip()}"
                     for i, pg in enumerate(reader.pages[:40], 1))


def _docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for ti, table in enumerate(doc.tables, 1):
        parts.append(f"[table {ti}]")
        for row in table.rows:
            parts.append("\t".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[sheet {ws.title}]")
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= 2000:
                parts.append("... (more rows not shown)")
                break
            parts.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def prepare_attachments(image_urls, file_urls, file_names):
    """Returns (file_text, media_urls, names).
    - documents/text files are read HERE (we already have the readers) and sent as text
    - images, and PDFs with no extractable text (scans), are sent as URLs for Gemini to look at directly"""
    media = [u for u in (image_urls or []) if isinstance(u, str) and u.startswith(("http://", "https://"))]
    blocks, names = [], []
    fnames = list(file_names or [])
    for i, url in enumerate(file_urls or []):
        name = fnames[i] if i < len(fnames) and fnames[i] else f"file_{i+1}"
        names.append(name)
        ext = _ext(url, name)
        if ext in _IMG_EXT:
            media.append(url)
            continue
        try:
            r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
            r.raise_for_status()
            data = r.content[:MAX_FILE_BYTES]
            if ext == ".pdf":
                text = _pdf_text(data)
                if len(re.sub(r"\[page \d+\]|\s", "", text)) < 20:      # scanned PDF: let Gemini read it natively
                    media.append(url)
                    continue
            elif ext == ".docx":
                text = _docx_text(data)
            elif ext in (".xlsx", ".xlsm"):
                text = _xlsx_text(data)
            else:
                text = data.decode("utf-8", errors="replace")
            blocks.append(f"--- FILE: {name} ---\n{text.strip()[:MAX_FILE_TEXT]}")
        except Exception as e:
            print(f"[LITE] couldn't read attachment {name!r}: {e}")
            blocks.append(f"--- FILE: {name} ---\n[Couldn't read this file.]")
    return "\n\n".join(blocks)[:MAX_FILE_TEXT * 2], media[:4], names


def _clean_history(history: Optional[List[dict]]) -> List[dict]:
    out = []
    for h in (history or [])[-HISTORY_MESSAGES:]:
        text = (h.get("content") or "")
        # drop the internal "[INTERNAL MEMORY NOTE ..." suffix (image/source URLs) — pure token waste here
        text = text.split("\n\n[INTERNAL MEMORY NOTE", 1)[0].strip()[:HISTORY_CHARS]
        if text:
            out.append({"role": "assistant" if h.get("role") == "assistant" else "user", "content": text})
    return out


def lite_events(question: str, history: Optional[List[dict]] = None, image_urls=None,
                file_urls=None, file_names=None) -> Iterator[dict]:
    file_text, media_urls, _names = prepare_attachments(image_urls, file_urls, file_names)
    payload = {"message": question, "history": _clean_history(history),
               "file_text": file_text, "media_urls": media_urls}
    headers = {"X-Service-Key": SERVICE_API_KEY}

    blocks: List[dict] = []     # ordered text/images blocks, same idea as the full agent's final event
    images: List[dict] = []
    answer = ""
    cur_text = ""
    started = False

    timeout = (10, 60) if (file_text or media_urls) else TIMEOUT
    try:
        resp = requests.post(LITE_STREAM_URL, json=payload, headers=headers, stream=True, timeout=timeout)
    except Exception as e:
        raise LiteUnavailable(f"connect failed: {e}")

    if resp.status_code != 200:
        resp.close()
        raise LiteUnavailable(f"service returned {resp.status_code}")

    try:
        for raw in resp.iter_lines(decode_unicode=True):
            if not raw:
                continue
            try:
                ev = json.loads(raw)
            except ValueError:
                continue
            kind = ev.get("type")

            if kind == "text":
                delta = ev.get("delta") or ""
                if delta:
                    started = True
                    answer += delta
                    cur_text += delta
                    yield {"type": "text", "delta": delta}

            elif kind == "images":
                imgs = ev.get("images") or []
                if imgs:
                    started = True
                    if cur_text:
                        blocks.append({"type": "text", "text": cur_text, "segments": build_answer_segments(cur_text)})
                        cur_text = ""
                    images.extend(imgs)
                    blocks.append({"type": "images", "images": imgs})
                    yield {"type": "block", "kind": "images", "images": imgs}

            elif kind == "error":
                print(f"[LITE] service reported: {ev.get('message')}")
            elif kind == "done":
                break
    except Exception as e:
        if not started:
            raise LiteUnavailable(f"stream failed before any output: {e}")
        print(f"[LITE] stream cut mid-answer, keeping partial: {e}")
    finally:
        resp.close()

    if not started:
        raise LiteUnavailable("service returned nothing")

    if cur_text:
        blocks.append({"type": "text", "text": cur_text, "segments": build_answer_segments(cur_text)})

    yield {
        "type": "final",
        "answer": answer,
        "segments": build_answer_segments(answer),
        "blocks": blocks,
        "suggestions": [],
        "sources": [],
        "images": images,
        "file": None,
        "files": [],
        "codes": [],
        "provider": "lite-gemini",
        "models": [],
        "model_count": 0,
    }
