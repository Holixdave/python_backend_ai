"""Offline tests for the workspace (no Firebase/Supabase/network needed).
Run:  python -m unittest discover -s tests -v"""
import io, os, shutil, sys, tempfile, time, types, unittest, zipfile
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
TMP = tempfile.mkdtemp(prefix="wstest-")
os.environ["WORKSPACE_ROOT"] = os.path.join(TMP, "ws")
os.environ["WORKSPACE_PERSIST"] = "1"

# stub gpt2_functions (it needs Firebase/Groq at import) with an in-memory "Supabase"
FAKE_BUCKET = {}
stub = types.ModuleType("gpt2_functions")
def _fake_upload(uid, name, data, ctype):
    FAKE_BUCKET[f"{uid}/{name}"] = data
    return f"https://fake/{uid}/{name}"
stub._upload_bytes_to_supabase = _fake_upload
sys.modules["gpt2_functions"] = stub

import gpt2_workspace as w
import gpt2_workspace_store as store

# fake Firestore doc + fake HTTP download for the store
FAKE_META = {}
class _Ref:
    def __init__(self, uid): self.uid = uid
    def get(self):
        d = FAKE_META.get(self.uid)
        return types.SimpleNamespace(exists=d is not None, to_dict=lambda: dict(d or {}))
    def set(self, data, merge=False):
        FAKE_META.setdefault(self.uid, {}).update(data) if merge else FAKE_META.__setitem__(self.uid, dict(data))
store._meta_ref = _Ref
class _Resp:
    def __init__(self, b): self.content = b
    def raise_for_status(self): pass
store.requests.get = lambda url, timeout=0: _Resp(FAKE_BUCKET[url.replace("https://fake/", "")])


class Safety(unittest.TestCase):
    def test_path_escape_blocked(self):
        for p in ["../x", "a/../../x"]:
            with self.assertRaises(ValueError):
                w.workspace_read("u1", p)

    def test_zip_slip_and_symlink_skipped(self):
        b = io.BytesIO()
        with zipfile.ZipFile(b, "w") as zf:
            zf.writestr("../../evil.txt", "x"); zf.writestr("ok/a.txt", "hi")
            i = zipfile.ZipInfo("ln"); i.external_attr = 0o120777 << 16; zf.writestr(i, "/etc/passwd")
        r = w._extract_zip_bytes("u1", b.getvalue(), "slip")
        self.assertTrue(r["success"]); self.assertEqual(r["files"], 1)
        self.assertFalse(os.path.exists(os.path.join(TMP, "evil.txt")))

    def test_users_are_separate(self):
        w.workspace_write("alice", "secret.txt", "A")
        self.assertFalse(w.workspace_read("bob", "secret.txt")["success"])


class EditAndRead(unittest.TestCase):
    def setUp(self):
        w.workspace_write("e1", "app.py", "a = 1\nb = 2\nb = 2\n")

    def test_edit_unique_and_ambiguous(self):
        self.assertTrue(w.workspace_edit("e1", "app.py", "a = 1", "a = 10")["success"])
        r = w.workspace_edit("e1", "app.py", "b = 2", "b = 3")
        self.assertFalse(r["success"]); self.assertIn("2 places", r["error"])
        self.assertTrue(w.workspace_edit("e1", "app.py", "b = 2", "b = 3", replace_all=True)["success"])
        self.assertIn("b = 3", w.workspace_read("e1", "app.py")["content"])

    def test_edit_miss_gives_hint(self):
        r = w.workspace_edit("e1", "app.py", "a = 1  # nope", "x")
        self.assertFalse(r["success"]); self.assertIn("not found", r["error"])

    def test_big_file_paging(self):
        w.workspace_write("e1", "big.txt", "\n".join(f"line {i}" for i in range(1, 1001)))
        r = w.workspace_read("e1", "big.txt", 1, 5)
        self.assertEqual(r["total_lines"], 1000); self.assertEqual(r["end_line"], 5)
        r2 = w.workspace_read("e1", "big.txt")
        self.assertIn("more", r2)

    def test_docx_and_xlsx_text(self):
        from docx import Document
        from openpyxl import Workbook
        root = w._user_root("e1")
        d = Document(); d.add_paragraph("hello docx world"); d.save(os.path.join(root, "t.docx"))
        wb = Workbook(); wb.active.append(["name", "score"]); wb.active.append(["ada", 7]); wb.save(os.path.join(root, "t.xlsx"))
        self.assertIn("hello docx world", w.workspace_read("e1", "t.docx")["content"])
        self.assertIn("ada\t7", w.workspace_read("e1", "t.xlsx")["content"])


class Persistence(unittest.TestCase):
    def test_survives_redeploy(self):
        uid = "persist-user"
        w.workspace_write(uid, "proj/main.py", "print('hi')")
        store.hydrate  # hydrated on first _user_root call
        self.assertTrue(store.persist_now(uid))
        self.assertIn(uid, FAKE_META); self.assertTrue(FAKE_META[uid]["snapshot_url"])
        # simulate Render wiping the disk + restarting the process
        shutil.rmtree(w._root_path(uid))
        store._hydrated.discard(uid)
        r = w.workspace_read(uid, "proj/main.py")
        self.assertTrue(r["success"], r); self.assertIn("print('hi')", r["content"])

    def test_failed_restore_never_overwrites_good_snapshot(self):
        uid = "safe-user"
        w.workspace_write(uid, "keep.txt", "precious"); store.persist_now(uid)
        good = dict(FAKE_META[uid])
        shutil.rmtree(w._root_path(uid)); store._hydrated.discard(uid)
        real = store.requests.get
        store.requests.get = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("network down"))
        try:
            w.workspace_write(uid, "new.txt", "x")           # triggers hydrate (fails) then a write
            self.assertFalse(store.persist_now(uid))          # must refuse to save
        finally:
            store.requests.get = real
        self.assertEqual(FAKE_META[uid]["snapshot_url"], good["snapshot_url"])

    def test_oversize_not_persisted(self):
        uid = "huge-user"; w.workspace_write(uid, "a.txt", "x")
        old = store.MAX_SNAPSHOT_BYTES; store.MAX_SNAPSHOT_BYTES = 10
        try:
            self.assertFalse(store.persist_now(uid))
            self.assertIn("NOT being saved", store.status(uid))
        finally:
            store.MAX_SNAPSHOT_BYTES = old


class RunInWorkspace(unittest.TestCase):
    def test_run_syncs_changes_and_reports_isolation(self):
        import gpt2_workspace_run as run
        w.workspace_write("r1", "proj/hello.py", "open('made.txt','w').write('ok'); print('ran')")
        res = run.workspace_run("r1", "python3 hello.py", cwd="proj")
        self.assertTrue(res["success"], res); self.assertIn("ran", res["stdout"])
        self.assertIn("proj/made.txt", res["changed_files"])
        self.assertEqual(w.workspace_read("r1", "proj/made.txt")["content"], "1: ok")

    def test_run_cannot_see_other_workspaces(self):
        import gpt2_workspace_run as run
        w.workspace_write("victim", "secret.txt", "TOP-SECRET")
        w.workspace_write("r2", "x.txt", "x")
        cmd = f"cat {w._root_path('victim')}/secret.txt; ls {w.WORKSPACE_ROOT}"
        res = run.workspace_run("r2", cmd)
        self.assertNotIn("TOP-SECRET", res["stdout"])
        if res["isolation"]["runs_as"] == "nobody":
            self.assertIn("denied", res["stderr"].lower())

    def test_timeout_kills(self):
        import gpt2_workspace_run as run
        w.workspace_write("r3", "x.txt", "x")
        t = time.time(); res = run.workspace_run("r3", "sleep 30", timeout_s=2)
        self.assertTrue(res["timed_out"]); self.assertLess(time.time() - t, 10)


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
