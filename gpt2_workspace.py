#!/usr/bin/env python3
# gpt2_workspace.py — a real, persistent WORKSPACE per user for the AI.
# ─────────────────────────────────────────────────────────────────────────────
# WHY THIS EXISTS
#   Before: the AI only had one-shot tools. fetch_github_file pulled ONE file
#   through the GitHub API and the result got cut at 20,000 chars; attached
#   .zip files were read as text (garbage); nothing could be explored later.
#
#   Now: the AI gets a folder on disk it can fill and explore across turns —
#   the same idea as Claude's own workspace:
#     workspace_github_clone  -> download the WHOLE repo once (no truncation)
#     workspace_unzip         -> open any .zip (uploaded or already in workspace)
#     workspace_list / workspace_grep / workspace_read (line ranges)
#                             -> look at only the part you need, any file size
#     workspace_write / workspace_delete / workspace_zip -> change + hand back
#
#   Truncation stops being a problem because the AI never has to load a whole
#   repo into one reply: it clones once, then greps and reads line ranges.
#
# SAFETY (all paths are per-user, nothing escapes the user's folder)
#   - every path is resolved and checked to stay inside the user's root
#   - zip-slip blocked, symlinks in archives skipped, zip-bomb caps enforced
#     while streaming (not just by trusting the header sizes)
#   - per-user total size cap + per-file cap + file-count cap
#
# PERSISTENCE: set WORKSPACE_ROOT to a mounted disk (Render paid-plan disk,
# a volume, etc.). The default (/tmp) is wiped on redeploy/restart.
# ─────────────────────────────────────────────────────────────────────────────

import hashlib
import io
import os
import re
import shutil
import stat
import threading
import time
import zipfile
from collections import defaultdict, deque
from typing import Optional

import requests

WORKSPACE_ROOT = os.environ.get("WORKSPACE_ROOT", "/tmp/ai_workspaces")
GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN")

MAX_USER_BYTES = 300 * 1024 * 1024        # whole workspace per user
MAX_FILE_BYTES = 25 * 1024 * 1024         # any single file
MAX_ARCHIVE_BYTES = 150 * 1024 * 1024     # total uncompressed from one archive
MAX_DOWNLOAD_BYTES = 80 * 1024 * 1024     # compressed zip we'll download
MAX_ARCHIVE_ENTRIES = 8000
MAX_LIST_ENTRIES = 300
MAX_READ_CHARS = 20_000
MAX_GREP_RESULTS = 60

MARKER = ".ws_hydrated"   # internal persistence marker, hidden from every listing
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", ".idea", ".dart_tool", "build", "dist"}
_BINARY_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".ico", ".pdf", ".zip", ".gz", ".db",
               ".sqlite", ".pyc", ".so", ".dll", ".exe", ".woff", ".woff2", ".ttf", ".mp3", ".mp4", ".wav"}
_REPO_RE = re.compile(r"^[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+$")


# ── paths ────────────────────────────────────────────────────────────────────
def _root_path(userid: Optional[str]) -> str:
    """Where this user's workspace lives (pure: no mkdir, no restore)."""
    raw = userid or "anonymous"
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", raw)[:48]
    # hash suffix so two different ids that sanitize the same never share a folder
    tag = hashlib.sha1(raw.encode()).hexdigest()[:8]
    return os.path.realpath(os.path.join(WORKSPACE_ROOT, f"{safe}-{tag}"))


def _user_root(userid: Optional[str]) -> str:
    root = _root_path(userid)
    # 0700: when the server runs as root, sandboxed code (which runs as
    # 'nobody') cannot even traverse into any workspace.
    os.makedirs(WORKSPACE_ROOT, mode=0o700, exist_ok=True)
    try:
        os.chmod(WORKSPACE_ROOT, 0o700)
    except OSError:
        pass
    os.makedirs(root, mode=0o700, exist_ok=True)
    try:
        from gpt2_workspace_store import hydrate  # restore from Supabase+Firebase after a redeploy
        hydrate(userid, root)
    except Exception as e:
        print(f"[WORKSPACE] hydrate skipped: {e}")
    return root


def _dirty(userid: Optional[str]) -> None:
    try:
        from gpt2_workspace_store import mark_dirty
        mark_dirty(userid)
    except Exception as e:
        print(f"[WORKSPACE] persist scheduling failed: {e}")


_rate_hits: dict = defaultdict(deque)
_rate_lock = threading.Lock()


def _rate_ok(userid: Optional[str], bucket: str, limit: int, per_s: int = 3600) -> bool:
    """Simple in-memory per-user limiter so one user can't hammer clone/run."""
    now = time.time()
    with _rate_lock:
        q = _rate_hits[(userid or "anonymous", bucket)]
        while q and now - q[0] > per_s:
            q.popleft()
        if len(q) >= limit:
            return False
        q.append(now)
        return True


def _safe_path(userid: Optional[str], rel: str = "") -> str:
    """Resolve `rel` inside the user's root; raise if it tries to escape."""
    root = _user_root(userid)
    rel = (rel or "").replace("\\", "/").lstrip("/")
    full = os.path.realpath(os.path.join(root, rel))
    if full != root and not full.startswith(root + os.sep):
        raise ValueError(f"Path '{rel}' is outside the workspace.")
    return full


def _rel(userid: Optional[str], full: str) -> str:
    return os.path.relpath(full, _user_root(userid)).replace(os.sep, "/")


def _dir_size(path: str) -> int:
    total = 0
    for dp, _dn, fn in os.walk(path):
        for f in fn:
            try:
                total += os.path.getsize(os.path.join(dp, f))
            except OSError:
                pass
    return total


def _is_probably_text(path: str) -> bool:
    if os.path.splitext(path)[1].lower() in _BINARY_EXT:
        return False
    try:
        with open(path, "rb") as f:
            return b"\0" not in f.read(2048)
    except OSError:
        return False


# ── list / tree ──────────────────────────────────────────────────────────────
def workspace_list(userid: str, path: str = "", depth: int = 2) -> dict:
    """
    Show what's in the workspace (or a folder inside it) as an indented tree
    with file sizes. depth = how many folder levels to expand (default 2).
    """
    base = _safe_path(userid, path)
    if not os.path.exists(base):
        return {"success": False, "error": f"'{path}' doesn't exist in the workspace."}
    if os.path.isfile(base):
        return {"success": True, "path": path, "type": "file", "size": os.path.getsize(base)}

    depth = max(1, min(int(depth or 2), 6))
    lines, count, truncated = [], 0, False

    def walk(d: str, level: int):
        nonlocal count, truncated
        try:
            entries = sorted(os.scandir(d), key=lambda e: (not e.is_dir(), e.name.lower()))
        except OSError:
            return
        for e in entries:
            if count >= MAX_LIST_ENTRIES:
                truncated = True
                return
            if e.is_symlink() or e.name == MARKER:
                continue
            indent = "  " * level
            if e.is_dir():
                if e.name in SKIP_DIRS:
                    lines.append(f"{indent}{e.name}/  (skipped)")
                    count += 1
                    continue
                lines.append(f"{indent}{e.name}/")
                count += 1
                if level + 1 < depth:
                    walk(e.path, level + 1)
            else:
                lines.append(f"{indent}{e.name}  ({e.stat().st_size} B)")
                count += 1

    walk(base, 0)
    out = {"success": True, "path": path or "/", "tree": "\n".join(lines) if lines else "(empty)"}
    if not path:
        try:
            from gpt2_workspace_store import status
            st = status(userid)
            if st:
                out["persistence"] = st
        except Exception:
            pass
    if truncated:
        out["note"] = (f"Listing capped at {MAX_LIST_ENTRIES} entries — "
                       "pass a deeper `path` to look inside a specific folder.")
    return out


# ── document text extraction (pdf / docx / xlsx) ────────────────────────────
_EXTRACTABLE = {".pdf", ".docx", ".xlsx", ".xlsm"}
_extract_cache: dict = {}


def _extract_document_text(full: str, ext: str) -> str:
    key = (full, os.path.getmtime(full))
    if key in _extract_cache:
        return _extract_cache[key]
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(full)
        text = "\n".join(f"[page {i}]\n{(pg.extract_text() or '').strip()}" for i, pg in enumerate(reader.pages, 1))
    elif ext == ".docx":
        from docx import Document
        doc = Document(full)
        parts = [p.text for p in doc.paragraphs]
        for ti, table in enumerate(doc.tables, 1):
            parts.append(f"[table {ti}]")
            for row in table.rows:
                parts.append("\t".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:  # xlsx / xlsm
        from openpyxl import load_workbook
        wb = load_workbook(full, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[sheet {ws.title}]")
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= 5000:
                    parts.append("... (more rows not shown)")
                    break
                parts.append("\t".join("" if v is None else str(v) for v in row))
        text = "\n".join(parts)
    if len(_extract_cache) > 20:
        _extract_cache.clear()
    _extract_cache[key] = text
    return text


# ── read ─────────────────────────────────────────────────────────────────────
def workspace_read(userid: str, path: str, start_line: int = 1, end_line: Optional[int] = None) -> dict:
    """
    Read a text file from the workspace by LINE RANGE. Returns numbered lines
    plus total_lines, so a big file is read in pieces (call again with the
    next range) instead of being truncated. Default range is ~the first 400
    lines; output is capped at ~20k chars and tells you where to continue.
    """
    full = _safe_path(userid, path)
    if not os.path.isfile(full):
        return {"success": False, "error": f"'{path}' is not a file in the workspace."}
    ext = os.path.splitext(full)[1].lower()
    if ext in _EXTRACTABLE:
        try:
            text = _extract_document_text(full, ext)
        except Exception as e:
            return {"success": False, "error": f"Couldn't read {ext} file: {e}"}
        if not text.strip():
            return {"success": False, "error": "No extractable text (it may be a scanned/image-only file)."}
        all_lines = text.splitlines()
    elif not _is_probably_text(full):
        return {"success": False, "error": f"'{path}' looks binary — can't show it as text."}
    else:
        with open(full, "r", encoding="utf-8", errors="replace") as f:
            all_lines = f.read().splitlines()
    total = len(all_lines)

    start = max(1, int(start_line or 1))
    end = int(end_line) if end_line else min(total, start + 399)
    end = min(end, total)
    if start > total and total > 0:
        return {"success": False, "error": f"start_line {start} is past the end ({total} lines).", "total_lines": total}

    out, used = [], 0
    last = start - 1
    for n in range(start, end + 1):
        line = f"{n}: {all_lines[n - 1]}"
        if used + len(line) + 1 > MAX_READ_CHARS:
            break
        out.append(line)
        used += len(line) + 1
        last = n

    res = {"success": True, "path": path, "start_line": start, "end_line": last,
           "total_lines": total, "content": "\n".join(out)}
    if last < total:
        res["more"] = f"Continue with start_line={last + 1} to read the rest ({total - last} lines left)."
    return res


# ── write / delete ───────────────────────────────────────────────────────────
def workspace_write(userid: str, path: str, content: str, deliver: bool = False) -> dict:
    """Create or overwrite a text file in the workspace (parent folders are created).
    Set deliver=true to ALSO send the file to the user as a download card right now —
    that is how you hand over a file you just wrote (one call, no output_files needed)."""
    if not path or path.endswith("/"):
        return {"success": False, "error": "Give a file path, e.g. 'src/app.py'."}
    data = (content or "").encode("utf-8")
    if len(data) > MAX_FILE_BYTES:
        return {"success": False, "error": "File too large."}
    full = _safe_path(userid, path)
    if full == _user_root(userid):
        return {"success": False, "error": "Can't overwrite the workspace root."}
    if _dir_size(_user_root(userid)) + len(data) > MAX_USER_BYTES:
        return {"success": False, "error": "Workspace is full — delete something first."}
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "wb") as f:
        f.write(data)
    _dirty(userid)
    out = {"success": True, "path": _rel(userid, full), "bytes": len(data)}
    if deliver:
        try:
            from gpt2_output import output_files   # lazy: gpt2_output imports this module
            res = output_files(userid, [out["path"]])
            out["delivered"] = res.get("files") or []
            if res.get("failed"):
                out["delivery_failed"] = res["failed"]
        except Exception as e:
            out["delivery_failed"] = [{"file": out["path"], "error": str(e)}]
    return out


def workspace_delete(userid: str, path: str) -> dict:
    """Delete a file or folder from the workspace."""
    full = _safe_path(userid, path)
    if full == _user_root(userid):
        return {"success": False, "error": "Refusing to delete the whole workspace; give a path inside it."}
    if not os.path.exists(full):
        return {"success": False, "error": f"'{path}' doesn't exist."}
    if os.path.isdir(full) and not os.path.islink(full):
        shutil.rmtree(full)
    else:
        os.remove(full)
    _dirty(userid)
    return {"success": True, "deleted": path}


# ── grep ─────────────────────────────────────────────────────────────────────
def workspace_grep(userid: str, pattern: str, path: str = "", case_sensitive: bool = False) -> dict:
    """
    Regex-search every text file under a workspace folder. Returns
    file:line: text matches (capped). Use this to FIND the right place in a
    big repo, then workspace_read that line range.
    """
    try:
        rx = re.compile(pattern, 0 if case_sensitive else re.IGNORECASE)
    except re.error as e:
        return {"success": False, "error": f"Bad regex: {e}"}

    base = _safe_path(userid, path)
    if not os.path.exists(base):
        return {"success": False, "error": f"'{path}' doesn't exist."}

    matches, truncated = [], False
    targets = [base] if os.path.isfile(base) else None
    if targets is None:
        targets = []
        for dp, dn, fn in os.walk(base):
            dn[:] = [d for d in dn if d not in SKIP_DIRS and not os.path.islink(os.path.join(dp, d))]
            targets += [os.path.join(dp, f) for f in sorted(fn)]

    for fp in targets:
        if os.path.islink(fp) or not _is_probably_text(fp):
            continue
        try:
            with open(fp, "r", encoding="utf-8", errors="replace") as f:
                for i, line in enumerate(f, 1):
                    if rx.search(line):
                        matches.append(f"{_rel(userid, fp)}:{i}: {line.rstrip()[:200]}")
                        if len(matches) >= MAX_GREP_RESULTS:
                            truncated = True
                            break
        except OSError:
            continue
        if truncated:
            break

    res = {"success": True, "matches": matches, "count": len(matches)}
    if truncated:
        res["note"] = f"Stopped at {MAX_GREP_RESULTS} matches — narrow the pattern or path."
    return res


# ── unzip ────────────────────────────────────────────────────────────────────
def _extract_zip_bytes(userid: str, data: bytes, dest: str, strip_top: bool = False) -> dict:
    """Shared safe extractor used by workspace_unzip, uploads and GitHub clone."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return {"success": False, "error": "That isn't a valid zip file."}

    dest_full = _safe_path(userid, dest)
    infos = zf.infolist()
    if len(infos) > MAX_ARCHIVE_ENTRIES:
        return {"success": False, "error": f"Archive has {len(infos)} entries (limit {MAX_ARCHIVE_ENTRIES})."}
    if sum(i.file_size for i in infos) > MAX_ARCHIVE_BYTES:
        return {"success": False, "error": "Archive is too large once unpacked."}
    if _dir_size(_user_root(userid)) + sum(i.file_size for i in infos) > MAX_USER_BYTES:
        return {"success": False, "error": "Not enough workspace space left for this archive."}

    # Strip a single shared top-level folder (GitHub zipballs and most
    # project zips wrap everything in one) so paths come out clean.
    prefix = ""
    if strip_top:
        tops = {i.filename.split("/", 1)[0] for i in infos if i.filename.strip("/")}
        if len(tops) == 1 and all("/" in i.filename or i.is_dir() for i in infos):
            prefix = tops.pop() + "/"

    written, skipped, total = 0, [], 0
    for info in infos:
        name = info.filename
        if prefix and name.startswith(prefix):
            name = name[len(prefix):]
        if not name or name.endswith("/"):
            continue
        # symlink entries (unix mode bits live in the high 16 bits)
        if stat.S_ISLNK(info.external_attr >> 16):
            skipped.append(f"{name} (symlink)")
            continue
        if info.file_size > MAX_FILE_BYTES:
            skipped.append(f"{name} (too big)")
            continue
        target = os.path.realpath(os.path.join(dest_full, name))
        if target != dest_full and not target.startswith(dest_full + os.sep):
            skipped.append(f"{name} (unsafe path)")
            continue
        os.makedirs(os.path.dirname(target), exist_ok=True)
        # stream with a hard cap: header sizes can lie
        with zf.open(info) as src, open(target, "wb") as out:
            copied = 0
            while True:
                chunk = src.read(65536)
                if not chunk:
                    break
                copied += len(chunk)
                total += len(chunk)
                if copied > MAX_FILE_BYTES or total > MAX_ARCHIVE_BYTES:
                    out.close()
                    os.remove(target)
                    return {"success": False, "error": "Archive expanded past the size limit — stopped."}
                out.write(chunk)
        written += 1

    _dirty(userid)
    res = {"success": True, "extracted_to": _rel(userid, dest_full) or "/", "files": written, "bytes": total}
    if skipped:
        res["skipped"] = skipped[:20]
    return res


def workspace_unzip(userid: str, path: str, dest: Optional[str] = None) -> dict:
    """
    Open a .zip that's already in the workspace and extract it. dest defaults
    to a folder named after the zip. After this, use workspace_list /
    workspace_grep / workspace_read on the extracted files.
    (Zips the user attaches to a message are saved + unzipped automatically.)
    """
    full = _safe_path(userid, path)
    if not os.path.isfile(full):
        return {"success": False, "error": f"'{path}' isn't in the workspace."}
    if os.path.getsize(full) > MAX_DOWNLOAD_BYTES:
        return {"success": False, "error": "Zip too large."}
    if not dest:
        dest = os.path.splitext(_rel(userid, full))[0]
    with open(full, "rb") as f:
        data = f.read()
    res = _extract_zip_bytes(userid, data, dest, strip_top=True)
    if res.get("success"):
        res["tree"] = workspace_list(userid, res["extracted_to"], depth=2).get("tree", "")
    return res


def workspace_ingest_zip_url(userid: str, url: str, filename: str) -> dict:
    """
    Backend-side (NOT an AI tool): called when the user ATTACHES a .zip.
    Downloads it, keeps the original in uploads/, unzips it into uploads/<name>/.
    """
    try:
        resp = requests.get(url, timeout=30, stream=True, headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
        buf = io.BytesIO()
        for chunk in resp.iter_content(65536):
            buf.write(chunk)
            if buf.tell() > MAX_DOWNLOAD_BYTES:
                return {"success": False, "error": "Zip is larger than the download limit."}
    except Exception as e:
        return {"success": False, "error": f"Couldn't download the zip: {e}"}

    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", os.path.basename(filename or "upload.zip")) or "upload.zip"
    if not safe_name.lower().endswith(".zip"):
        safe_name += ".zip"
    saved = workspace_write_bytes(userid, f"{UPLOADS_DIR}/{safe_name}", buf.getvalue())
    if not saved.get("success"):
        return saved
    folder = f"{UPLOADS_DIR}/{os.path.splitext(safe_name)[0]}"
    res = _extract_zip_bytes(userid, buf.getvalue(), folder, strip_top=True)
    res["zip_saved_as"] = f"{UPLOADS_DIR}/{safe_name}"
    if res.get("success"):
        res["tree"] = workspace_list(userid, res["extracted_to"], depth=2).get("tree", "")
    return res


def workspace_write_bytes(userid: str, path: str, data: bytes) -> dict:
    full = _safe_path(userid, path)
    if _dir_size(_user_root(userid)) + len(data) > MAX_USER_BYTES:
        return {"success": False, "error": "Workspace is full — delete something first."}
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "wb") as f:
        f.write(data)
    _dirty(userid)
    return {"success": True, "path": _rel(userid, full)}


# ── GitHub ───────────────────────────────────────────────────────────────────
def _parse_repo(repo: str) -> Optional[str]:
    repo = (repo or "").strip()
    m = re.match(r"^(?:https?://)?(?:www\.)?github\.com/([^/\s]+)/([^/\s#?]+)", repo)
    if m:
        repo = f"{m.group(1)}/{m.group(2)}"
    repo = re.sub(r"\.git$", "", repo)
    return repo if _REPO_RE.match(repo) else None


def workspace_github_clone(userid: str, repo: str, ref: Optional[str] = None, dest: Optional[str] = None) -> dict:
    """
    Download an ENTIRE GitHub repo into the workspace — all files, no
    per-file truncation. repo: "owner/name" or a github.com URL. ref:
    branch/tag/commit (default = repo's default branch). Then explore with
    workspace_list / workspace_grep / workspace_read. Private repos work if
    the server has a GITHUB_TOKEN with access.
    """
    if not _rate_ok(userid, "clone", 10):
        return {"success": False, "error": "Clone limit reached (10 per hour) — try again later."}
    slug = _parse_repo(repo)
    if not slug:
        return {"success": False, "error": "repo must look like 'owner/name' or a github.com URL."}
    if ref and not re.match(r"^[A-Za-z0-9._/-]{1,100}$", ref):
        return {"success": False, "error": "Invalid ref."}

    headers = {"Accept": "application/vnd.github+json", "User-Agent": "workspace-bot"}
    if GITHUB_TOKEN:
        headers["Authorization"] = f"Bearer {GITHUB_TOKEN}"
    api = f"https://api.github.com/repos/{slug}/zipball" + (f"/{ref}" if ref else "")

    try:
        resp = requests.get(api, headers=headers, timeout=60, stream=True, allow_redirects=True)
        if resp.status_code == 404:
            return {"success": False, "error": f"{slug} not found (or private and the server token can't see it)."}
        if resp.status_code in (401, 403):
            return {"success": False, "error": f"GitHub refused access ({resp.status_code}) — likely rate limit or a private repo."}
        resp.raise_for_status()
        buf = io.BytesIO()
        for chunk in resp.iter_content(65536):
            buf.write(chunk)
            if buf.tell() > MAX_DOWNLOAD_BYTES:
                return {"success": False, "error": "Repo is bigger than the download limit; fetch specific files instead."}
    except Exception as e:
        return {"success": False, "error": f"GitHub download failed: {e}"}

    dest = dest or slug.split("/")[1]
    res = _extract_zip_bytes(userid, buf.getvalue(), dest, strip_top=True)
    if res.get("success"):
        res["repo"] = slug
        res["ref"] = ref or "default branch"
        res["tree"] = workspace_list(userid, res["extracted_to"], depth=2).get("tree", "")
    return res


# ── hand results back ────────────────────────────────────────────────────────
def workspace_zip(userid: str, path: str = "", zip_name: str = "workspace.zip") -> dict:
    """
    Pack a workspace folder (or the whole workspace) into one .zip and upload
    it, returning a download URL for the user.
    """
    from gpt2_functions import _upload_bytes_to_supabase  # lazy: avoid import cycle

    base = _safe_path(userid, path)
    if not os.path.isdir(base):
        return {"success": False, "error": f"'{path}' isn't a folder in the workspace."}
    if not zip_name.lower().endswith(".zip"):
        zip_name += ".zip"
    zip_name = re.sub(r"[^A-Za-z0-9._-]", "_", zip_name)

    buf = io.BytesIO()
    count = 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for dp, dn, fn in os.walk(base):
            dn[:] = [d for d in dn if d not in SKIP_DIRS and not os.path.islink(os.path.join(dp, d))]
            for f in fn:
                fp = os.path.join(dp, f)
                if os.path.islink(fp) or f == MARKER:
                    continue
                zf.write(fp, os.path.relpath(fp, base))
                count += 1
    data = buf.getvalue()
    if count == 0:
        return {"success": False, "error": "Nothing to zip."}
    if len(data) > 10 * 1024 * 1024:
        return {"success": False, "error": "Zip would be over the 10MB upload limit — zip a smaller folder."}
    url = _upload_bytes_to_supabase(userid or "anonymous", zip_name, data, "application/zip")
    if not url:
        return {"success": False, "error": "Upload failed."}
    return {"success": True, "filename": zip_name, "url": url, "files": count, "size": len(data)}


# ── surgical edit ────────────────────────────────────────────────────────────
def workspace_edit(userid: str, path: str, find_text: str, replace_text: str, replace_all: bool = False) -> dict:
    """
    Edit part of a workspace text file by exact find/replace — do NOT retype the
    whole file. find_text must match exactly ONCE (or set replace_all=true).
    On a miss or multiple matches the error tells you where, so read those
    lines with workspace_read and retry with a longer, unique find_text.
    """
    full = _safe_path(userid, path)
    if not os.path.isfile(full):
        return {"success": False, "error": f"'{path}' is not a file in the workspace."}
    if not find_text:
        return {"success": False, "error": "find_text can't be empty."}
    try:
        with open(full, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        return {"success": False, "error": "File isn't valid UTF-8 text."}

    n = text.count(find_text)
    if n == 0:
        first = next((ln.strip() for ln in find_text.splitlines() if ln.strip()), "")
        hint = ""
        if first:
            for i, ln in enumerate(text.splitlines(), 1):
                if first in ln:
                    hint = f" A line containing the first line of your find_text is at line {i} — read it and copy the exact text."
                    break
        return {"success": False, "error": "find_text not found (whitespace/indentation must match exactly)." + hint}
    if n > 1 and not replace_all:
        lines = []
        pos = 0
        for _ in range(min(n, 6)):
            idx = text.index(find_text, pos)
            lines.append(text[:idx].count("\n") + 1)
            pos = idx + max(1, len(find_text))
        return {"success": False, "error": f"find_text matches {n} places (lines {lines}). Add surrounding lines to make it unique, or pass replace_all=true."}

    first_line = text[:text.index(find_text)].count("\n") + 1
    new_text = text.replace(find_text, replace_text) if replace_all else text.replace(find_text, replace_text, 1)
    data = new_text.encode("utf-8")
    if len(data) > MAX_FILE_BYTES:
        return {"success": False, "error": "Result would be too large."}
    with open(full, "wb") as f:
        f.write(data)
    _dirty(userid)
    return {"success": True, "path": _rel(userid, full), "replacements": n if replace_all else 1,
            "first_change_at_line": first_line}


# ── edit by plain-English intent ─────────────────────────────────────────────
MAX_INTENT_EDIT_CHARS = 60_000   # bigger files: use workspace_read + workspace_edit instead


def workspace_edit_by_intent(userid: str, path: str, instruction: str) -> dict:
    """
    Edit ONE workspace text file by describing the change in plain English,
    e.g. "change the submit button color to blue" — you do NOT need the exact
    current text. A separate step reads the real file, works out the precise
    unique find/replace (retrying itself up to 3 times if the first guess
    isn't unique), then applies it with workspace_edit. Prefer this over
    workspace_edit when you can't see the file's exact wording. For files
    over ~60,000 characters, workspace_read the part you need and use
    workspace_edit with the exact text instead.
    """
    if not path or not str(path).strip():
        return {"success": False, "error": "no path given"}
    if not instruction or not str(instruction).strip():
        return {"success": False, "error": "no instruction given"}
    if not _rate_ok(userid, "intent_edit", 60):
        return {"success": False, "error": "Edit limit reached (60 per hour) — try again later."}
    try:
        full = _safe_path(userid, path)
    except ValueError as e:
        return {"success": False, "error": str(e)}
    if not os.path.isfile(full):
        return {"success": False, "error": f"'{path}' is not a file in the workspace. "
                                           "Use workspace_list to see what exists (uploads are in uploads/)."}
    try:
        with open(full, "r", encoding="utf-8") as f:
            content = f.read()
    except UnicodeDecodeError:
        return {"success": False, "error": "File isn't valid UTF-8 text."}
    if not content.strip():
        return {"success": False, "error": "file is empty, nothing to edit"}
    if len(content) > MAX_INTENT_EDIT_CHARS:
        return {"success": False, "error": f"File is {len(content)} characters (limit {MAX_INTENT_EDIT_CHARS}). "
                                           "Use workspace_read for the part you need, then workspace_edit."}

    try:  # lazy: gpt2_functions pulls in Firebase/provider keys
        from gpt2_functions import (
            _call_provider_chain_full, TEXT_PROVIDERS, _edit_executor_system_prompt,
            _FENCE_RE, EDIT_EXECUTOR_MAX_ATTEMPTS,
        )
    except Exception as e:
        return {"success": False, "error": f"editor model not available: {e}"}

    messages = [
        {"role": "system", "content": _edit_executor_system_prompt()},
        {"role": "user", "content": f"FILE CONTENT ({path}):\n\n{content}\n\nINSTRUCTION: {str(instruction).strip()}"},
    ]
    feedback = None
    for attempt in range(1, EDIT_EXECUTOR_MAX_ATTEMPTS + 1):
        if feedback:
            messages.append({"role": "user", "content": feedback})
        raw, _provider, _finish = _call_provider_chain_full(TEXT_PROVIDERS, messages, temperature=0.1, max_tokens=4000)
        if raw is None:
            return {"success": False, "error": "editor model unavailable — try again shortly"}
        messages.append({"role": "assistant", "content": raw})
        try:
            import json as _json
            parsed = _json.loads(_FENCE_RE.sub("", raw).strip())
            find_text = parsed.get("find", "") or ""
            replace_text = parsed.get("replace", "") or ""
        except Exception:
            feedback = ('That was not valid JSON. Respond with ONLY the JSON object '
                        '{"find": "...", "replace": "..."}, nothing else — no explanation, no markdown fences.')
            continue
        if not find_text:
            feedback = '"find" was empty. It must be a real, non-empty snippet copied exactly from the FILE CONTENT above.'
            continue
        n = content.count(find_text)
        if n == 0:
            feedback = ("That exact text was not found in the file — copy the snippet character-for-character "
                        "from the FILE CONTENT above (whitespace, quotes, indentation) and try again.")
            continue
        if n > 1:
            feedback = (f"That text matched {n} places in the file — it must match exactly once. Include a bit "
                        "more surrounding context (the line above or below) so it is unique, then try again.")
            continue
        result = workspace_edit(userid, path, find_text, replace_text)
        if result.get("success"):
            result["attempts"] = attempt
            result["instruction"] = str(instruction).strip()[:120]
        return result

    return {"success": False,
            "error": f"Couldn't land a unique edit after {EDIT_EXECUTOR_MAX_ATTEMPTS} attempts — "
                     "give a more specific instruction, or use workspace_read + workspace_edit with exact text."}


# ── uploads: ONE place for everything the user attaches ──────────────────────
# Attachments land in  <workspace>/uploads/  . The AI looks there (workspace_list),
# reads what it needs, and copies the files it wants to keep working on into the
# main workspace with workspace_copy. There is no second file store.
UPLOADS_DIR = "uploads"


def _safe_upload_name(name: str, fallback: str = "upload") -> str:
    base = os.path.basename((name or "").replace("\\", "/")) or fallback
    return re.sub(r"[^A-Za-z0-9._ -]", "_", base)[:120] or fallback


def workspace_save_upload(userid: Optional[str], filename: str, data: bytes) -> dict:
    """Backend-side (NOT an AI tool): store one attachment as uploads/<filename>.
    If the name is taken, adds a numeric suffix so nothing is overwritten."""
    name = _safe_upload_name(filename)
    stem, ext = os.path.splitext(name)
    n, rel = 1, f"{UPLOADS_DIR}/{name}"
    while os.path.exists(_safe_path(userid, rel)):
        n += 1
        rel = f"{UPLOADS_DIR}/{stem}_{n}{ext}"
    return workspace_write_bytes(userid, rel, data)


def workspace_read_text(userid: Optional[str], path: str):
    """Backend-side helper: (text, None) for a UTF-8 workspace file, else (None, error)."""
    try:
        full = _safe_path(userid, path)
    except ValueError as e:
        return None, str(e)
    if not os.path.isfile(full):
        return None, (f"'{path}' is not a file in the workspace. Use workspace_list to see what "
                      f"exists (the user's attachments are in {UPLOADS_DIR}/).")
    try:
        with open(full, "r", encoding="utf-8") as f:
            return f.read(), None
    except UnicodeDecodeError:
        return None, f"'{path}' is a binary file — it can't be edited as text."


def workspace_copy(userid: str, src: str, dest: Optional[str] = None) -> dict:
    """Copy a file or folder inside the workspace — mainly from uploads/ into the main
    workspace so you can edit/run it there (the original stays in uploads/).
    dest defaults to the same name at the top level of the workspace."""
    import shutil
    try:
        s_full = _safe_path(userid, src)
        d_rel = dest or os.path.basename(src.rstrip("/\\")) or "copy"
        d_full = _safe_path(userid, d_rel)
    except ValueError as e:
        return {"success": False, "error": str(e)}
    if not os.path.exists(s_full):
        return {"success": False, "error": f"'{src}' doesn't exist. Use workspace_list {UPLOADS_DIR} to see the user's uploads."}
    if os.path.exists(d_full):
        return {"success": False, "error": f"'{d_rel}' already exists — pick another dest or workspace_delete it first."}
    try:
        os.makedirs(os.path.dirname(d_full), exist_ok=True)
        if os.path.isdir(s_full):
            if _dir_size(s_full) + _dir_size(_user_root(userid)) > MAX_USER_BYTES:
                return {"success": False, "error": "Workspace would be over its size limit."}
            shutil.copytree(s_full, d_full, symlinks=False)
        else:
            if os.path.getsize(s_full) + _dir_size(_user_root(userid)) > MAX_USER_BYTES:
                return {"success": False, "error": "Workspace would be over its size limit."}
            shutil.copy2(s_full, d_full)
    except Exception as e:
        return {"success": False, "error": f"Copy failed: {e}"}
    _dirty(userid)
    return {"success": True, "copied": _rel(userid, s_full), "to": _rel(userid, d_full)}
